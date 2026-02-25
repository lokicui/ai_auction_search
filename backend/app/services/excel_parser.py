import uuid
import openpyxl
from io import BytesIO
from typing import Optional
from app.models.schemas import AssetCreate

HEADER_MAPPINGS = {
    "title": ["标题", "名称", "资产名称", "title"],
    "asset_type": ["类别", "分类", "类型", "资产类别", "资产类型", "category", "type"],
    "region": ["所在地", "地址", "位置", "城市", "地区", "区域", "location", "region"],
    "price": ["起拍价", "价格", "起始价", "底价", "price"],
    "area": ["面积", "建筑面积", "area"],
    "image_url": ["图片", "图片链接", "image", "image_url"],
    "description": ["描述", "详情", "资产描述", "说明", "description"],
    "status": ["状态", "拍卖状态", "status"],
}


def _match_header(text: str) -> str:
    t = text.strip().lower()
    for field, keywords in HEADER_MAPPINGS.items():
        for kw in keywords:
            if kw in t:
                return field
    return ""


def parse_excel(content: bytes) -> list[tuple[str, AssetCreate]]:
    """Parse Excel file and return list of (id, AssetCreate) tuples."""
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb.active

    headers = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col_map: dict[str, int] = {}
    for i, h in enumerate(headers):
        field = _match_header(h)
        if field and field not in col_map:
            col_map[field] = i

    if "title" not in col_map:
        col_map["title"] = 0

    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None for c in row):
            continue

        def _get(field: str, default=""):
            idx = col_map.get(field)
            if idx is not None and idx < len(row) and row[idx] is not None:
                return row[idx]
            return default

        title = str(_get("title", "")).strip()
        if not title:
            continue

        try:
            price = float(_get("price", 0))
        except (ValueError, TypeError):
            price = 0.0

        try:
            area = float(_get("area", 0))
        except (ValueError, TypeError):
            area = 0.0

        asset = AssetCreate(
            title=title,
            asset_type=str(_get("asset_type", "")),
            region=str(_get("region", "")),
            price=price,
            area=area,
            image_url=str(_get("image_url", "")),
            description=str(_get("description", "")),
            status=str(_get("status", "即将开拍")),
        )
        results.append((str(uuid.uuid4()), asset))

    wb.close()
    return results


def create_template() -> bytes:
    """Create an Excel template for asset import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资产导入模板"
    headers = ["标题", "描述", "类别", "所在地", "面积", "起拍价", "图片链接", "状态"]
    ws.append(headers)
    ws.append(["示例：杭州市西湖区三居室住宅", "精装修住宅，120㎡", "住宅", "浙江省杭州市西湖区", 120, 3500000, "", "即将开拍"])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
