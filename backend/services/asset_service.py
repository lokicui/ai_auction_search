import uuid
import openpyxl
from io import BytesIO
from typing import List
from models import Asset
from services.vector_service import VectorService


HEADER_MAPPINGS = {
    "title": ["标题", "名称", "资产名称", "title"],
    "description": ["描述", "详情", "资产描述", "说明", "description"],
    "category": ["类别", "分类", "类型", "资产类别", "category"],
    "location": ["所在地", "地址", "位置", "城市", "地区", "location"],
    "starting_price": ["起拍价", "价格", "起始价", "底价", "price"],
    "status": ["状态", "拍卖状态", "status"],
    "source_url": ["链接", "url", "网址", "来源"],
}


def _match_header(header_text: str) -> str:
    """Match a header text to a known field name."""
    h = header_text.strip().lower()
    for field, keywords in HEADER_MAPPINGS.items():
        for kw in keywords:
            if kw in h:
                return field
    return ""


class AssetService:
    def __init__(self, vector_service: VectorService):
        self.vector_service = vector_service

    def parse_excel(self, file_content: bytes) -> List[Asset]:
        wb = openpyxl.load_workbook(BytesIO(file_content), read_only=True)
        ws = wb.active

        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers.append(str(cell.value or ""))

        col_map: dict[str, int] = {}
        for i, h in enumerate(headers):
            field = _match_header(h)
            if field and field not in col_map:
                col_map[field] = i

        if "title" not in col_map:
            col_map["title"] = 0

        assets = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(cell is None for cell in row):
                continue

            def _get(field: str, default=""):
                idx = col_map.get(field)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    return row[idx]
                return default

            title = str(_get("title", "")).strip()
            if not title:
                continue

            try:
                price = float(_get("starting_price", 0))
            except (ValueError, TypeError):
                price = 0.0

            asset = Asset(
                id=str(uuid.uuid4()),
                title=title,
                description=str(_get("description", "")),
                category=str(_get("category", "")),
                location=str(_get("location", "")),
                starting_price=price,
                status=str(_get("status", "在售")),
                source_url=str(_get("source_url", "")),
            )
            assets.append(asset)

        wb.close()
        return assets

    def import_assets(self, assets: List[Asset]) -> int:
        if not assets:
            return 0

        ids = [a.id for a in assets]
        documents = [
            f"{a.title} {a.description} {a.category} {a.location}" for a in assets
        ]
        metadatas = []
        for a in assets:
            m = {
                "title": a.title,
                "description": a.description,
                "category": a.category,
                "location": a.location,
                "starting_price": a.starting_price,
                "status": a.status,
                "source_url": a.source_url,
            }
            metadatas.append(m)

        self.vector_service.add_assets(ids, documents, metadatas)
        return len(assets)

    def get_assets(self, page: int = 1, page_size: int = 20) -> tuple[List[Asset], int]:
        total = self.vector_service.count()
        offset = (page - 1) * page_size
        if offset >= total:
            return [], total

        result = self.vector_service.get_all(limit=page_size, offset=offset)
        assets = []
        for i, aid in enumerate(result["ids"]):
            meta = result["metadatas"][i]
            assets.append(
                Asset(
                    id=aid,
                    title=meta.get("title", ""),
                    description=meta.get("description", ""),
                    category=meta.get("category", ""),
                    location=meta.get("location", ""),
                    starting_price=float(meta.get("starting_price", 0)),
                    status=meta.get("status", "在售"),
                    source_url=meta.get("source_url", ""),
                )
            )
        return assets, total

    def delete_asset(self, asset_id: str):
        self.vector_service.delete([asset_id])

    def delete_assets(self, asset_ids: List[str]):
        self.vector_service.delete(asset_ids)
